from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.organizations.models import Bank, OkedCode

# Официальный справочник БИК Национального Банка Республики Казахстан:
# https://nationalbank.kz/ru/page/spravochnik-bik-rk-ps
BANKS = (
    ("ATYNKZKA", "949", "АО «Altyn Bank»", "«Altyn Bank» АҚ"),
    ("BKCHKZKA", "913", "АО ДБ «БАНК КИТАЯ В КАЗАХСТАНЕ»", "«ҚАЗАҚСТАНДАҒЫ ҚЫТАЙ БАНКІ» ЕБ АҚ"),
    ("BRKEKZKA", "914", "АО «Bereke Bank»", "«Bereke Bank» АҚ"),
    ("CASPKZKA", "722", "АО «KASPI BANK»", "«Kaspi Bank» АҚ"),
    (
        "CEDUKZKA",
        "766",
        "АО «Центральный депозитарий ценных бумаг»",
        "«Бағалы қағаздар орталық депозитарийі» АҚ",
    ),
    ("CITIKZKA", "832", "АО «Ситибанк Казахстан»", "«Ситибанк Қазақстан» АҚ"),
    ("DVKAKZKA", "907", "АО «Банк Развития Казахстана»", "«Қазақстанның Даму Банкі» АҚ"),
    ("EABRKZKA", "700", "Евразийский банк развития", "Еуразия даму банкі"),
    ("EURIKZKA", "948", "АО «Евразийский Банк»", "«Еуразиялық банк» АҚ"),
    (
        "GCVPKZ2A",
        "009",
        "НАО «Государственная корпорация «Правительство для граждан»",
        "«Азаматтарға арналған үкімет» мемлекеттік корпорациясы КЕАҚ",
    ),
    ("HCSKKZKA", "972", "АО «Отбасы банк»", "«Отбасы банк» АҚ"),
    ("HLALKZKZ", "246", "АО «Исламский Банк ADCB»", "«ADCB Ислам Банкі» АҚ"),
    ("HSBKKZKX", "601", "АО «Народный Банк Казахстана»", "«Қазақстан Халық Банкі» АҚ"),
    (
        "ICBKKZKX",
        "930",
        "АО «Торгово-промышленный Банк Китая в г. Алматы»",
        "«Алматы қаласындағы Қытай сауда-өнеркәсіп Банкі» АҚ",
    ),
    ("INEARUMM", "550", "Межгосударственный Банк", "Мемлекетаралық Банк"),
    ("INLMKZKA", "886", "АО «Home Credit Bank»", "«Home Credit Bank» АҚ"),
    ("IRTYKZKA", "965", "АО «ForteBank»", "«ForteBank» АҚ"),
    ("KCCJKZKK", "715", "АО «Клиринговый центр KASE»", "«KASE клиринг орталығы» АҚ"),
    ("KCJBKZKX", "856", "АО «Банк ЦентрКредит»", "«Банк ЦентрКредит» АҚ"),
    ("KICEKZKX", "927", "АО «Казахстанская фондовая биржа»", "«Қазақстан қор биржасы» АҚ"),
    ("KINCKZKA", "821", "АО «Банк Bank RBK»", "«Bank RBK Банкі» АҚ"),
    (
        "KKMFKZ2A",
        "070",
        "РГУ «Комитет государственного казначейства МФ РК»",
        "«ҚР ҚМ Мемлекеттік қазынашылық комитеті» РММ",
    ),
    ("KMFBKZKK", "719", "АО «KMF Банк»", "«KMF Банк» АҚ"),
    ("KPSTKZKA", "563", "АО «КАЗПОЧТА»", "«Қазпошта» АҚ"),
    ("KSNVKZKA", "551", "АО «Фридом Банк Казахстан»", "«Фридом Банк Қазақстан» АҚ"),
    (
        "KZIBKZKA",
        "885",
        "АО ДБ «КАЗАХСТАН-ЗИРААТ ИНТЕРНЕШНЛ БАНК»",
        "«ҚАЗАҚСТАН-ЗИРААТ ХАЛЫҚАРАЛЫҚ БАНКІ» ЕБ АҚ",
    ),
    ("MOKFKZKA", "724", "АО «Коммерческий Банк БиЭнКей»", "«БиЭнКей Коммерциялық Банкі» АҚ"),
    (
        "NBRKKZKX",
        "125",
        "РГУ «Национальный Банк Республики Казахстан»",
        "«Қазақстан Республикасының Ұлттық Банкі» РММ",
    ),
    ("NURSKZKX", "849", "АО «Нурбанк»", "«Нұрбанк» АҚ"),
    ("SHBKKZKA", "435", "АО «Шинхан Банк Казахстан»", "«Шинхан Банк Қазақстан» АҚ"),
    ("TSESKZKA", "998", "АО «Alatau City Bank»", "«Alatau City Bank» АҚ"),
    ("VTBAKZKZ", "432", "ДО АО «Банк ВТБ (Казахстан)»", "«Банк ВТБ (Қазақстан)» АҚ ЕҰ"),
    ("ZAJSKZ22", "896", "АО «Исламский банк «Заман-Банк»", "«Заман-Банк» Ислам банкі» АҚ"),
)


class Command(BaseCommand):
    help = "Идемпотентно загружает официальные справочники ОКЭД и банков."

    @transaction.atomic
    def handle(self, *args, **options):
        oked_path = Path(settings.BASE_DIR) / "apps" / "organizations" / "data" / "oked.xlsx"
        if not oked_path.exists():
            raise CommandError(f"Файл справочника ОКЭД не найден: {oked_path}")

        workbook = load_workbook(oked_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = {}
        for row in sheet.iter_rows(min_row=4, values_only=True):
            code = str(row[0] or "").strip()
            if len(code) == 5 and code.isdigit():
                rows[code] = {
                    "name_kk": str(row[1] or "").strip(),
                    "name_ru": str(row[2] or "").strip(),
                }
        workbook.close()

        existing = OkedCode.objects.in_bulk(rows, field_name="code")
        created = [
            OkedCode(code=code, **names) for code, names in rows.items() if code not in existing
        ]
        changed = []
        for code, item in existing.items():
            names = rows[code]
            if item.name_ru != names["name_ru"] or item.name_kk != names["name_kk"]:
                item.name_ru = names["name_ru"]
                item.name_kk = names["name_kk"]
                item.is_active = True
                changed.append(item)
        OkedCode.objects.bulk_create(created, batch_size=250)
        OkedCode.objects.bulk_update(changed, ["name_ru", "name_kk", "is_active"], batch_size=250)
        OkedCode.objects.exclude(code__in=rows).update(is_active=False)

        for bic, bank_code, name_ru, name_kk in BANKS:
            Bank.objects.update_or_create(
                bic=bic,
                defaults={
                    "bank_code": bank_code,
                    "name_ru": name_ru,
                    "name_kk": name_kk,
                    "is_active": True,
                },
            )

        self.stdout.write(
            self.style.SUCCESS(f"Справочники загружены: ОКЭД — {len(rows)}, банки — {len(BANKS)}.")
        )
